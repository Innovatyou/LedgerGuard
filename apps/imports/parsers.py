"""
Safe readers for CSV and XLSX source files.

Workbooks are opened with openpyxl in read-only, values-only mode: formulas are
not evaluated, macros are never executed, and macro-enabled packages
(containing ``vbaProject.bin``) are rejected outright.
"""

import csv
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024


class FileRejected(ValueError):
    pass


def inspect_file(data: bytes, extension: str, max_bytes: int):
    if len(data) == 0:
        raise FileRejected("The file is empty.")
    if len(data) > max_bytes:
        raise FileRejected(f"The file exceeds the maximum upload size of {max_bytes // (1024 * 1024)} MB.")
    if extension == ".xlsx":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise FileRejected("The file is not a valid .xlsx workbook.")
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            if "[Content_Types].xml" not in names:
                raise FileRejected("The file is not a valid .xlsx workbook.")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise FileRejected("Macro-enabled workbooks are not accepted.")
            if sum(info.file_size for info in zf.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise FileRejected("The workbook expands beyond the permitted size.")
    elif extension == ".csv":
        if b"\x00" in data[:4096]:
            raise FileRejected("The file does not look like a text CSV file.")
    else:
        raise FileRejected(f"Unsupported file type {extension}. Accepted: .csv, .xlsx")


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        text = repr(value)
        return text[:-2] if text.endswith(".0") else text
    if isinstance(value, int | Decimal):
        return str(value)
    return str(value).strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileRejected("Unable to decode the CSV file (expected UTF-8 or Windows-1252).")


def _unique_headers(headers):
    seen, result = {}, []
    for idx, header in enumerate(headers):
        name = (header or "").strip() or f"Column {idx + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        result.append(name)
    return result


def read_table(data: bytes, extension: str, header_row: int = 1, sheet_name: str = "", max_rows: int | None = None):
    """
    Return ``(headers, rows)`` where rows is a list of ``(row_number, {header: text})``.
    ``row_number`` is the 1-based row in the source file. Completely blank rows are skipped.
    """
    if extension == ".csv":
        text = _decode(data)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
        except csv.Error:
            dialect = csv.excel
        raw_rows = list(csv.reader(io.StringIO(text), dialect))
    elif extension == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    raise FileRejected(f"Worksheet {sheet_name!r} not found. Available: {', '.join(wb.sheetnames)}")
                ws = wb[sheet_name]
            else:
                ws = wb.worksheets[0]
            raw_rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        raise FileRejected(f"Unsupported file type {extension}")

    if len(raw_rows) < header_row:
        raise FileRejected("The file has no header row at the configured position.")
    headers = _unique_headers([str(h) if h is not None else "" for h in raw_rows[header_row - 1]])
    rows = []
    for offset, values in enumerate(raw_rows[header_row:], start=header_row + 1):
        if not any(v not in (None, "") for v in values):
            continue
        record = {headers[i]: (values[i] if i < len(values) else None) for i in range(len(headers))}
        rows.append((offset, record))
        if max_rows and len(rows) > max_rows:
            raise FileRejected(f"The file has more than {max_rows} data rows.")
    return headers, rows
