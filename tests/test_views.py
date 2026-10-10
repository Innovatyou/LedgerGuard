"""
UI tests: every page renders for authorised users on realistic demo data,
and the main workflows work end to end through the web forms.
"""

import io
from decimal import Decimal

import pytest
from django.core.management import call_command
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounting.models import Account, Journal
from apps.accounts.models import User
from apps.imports.models import ImportBatch, ImportProfile
from apps.organizations.models import Institution
from apps.reconciliation.models import MatchGroup, Reconciliation
from apps.reconciliation.services import build_statement
from apps.reporting.exports import safe_text
from apps.suspense.models import SuspenseItem

pytestmark = pytest.mark.django_db
PASSWORD = "Demo-Passw0rd-2026"


@pytest.fixture
def demo(db):
    call_command("seed_demo", password=PASSWORD, stdout=io.StringIO())
    inst = Institution.objects.get(code="demo-mfb")
    return inst, {u.username: u for u in User.objects.filter(memberships__institution=inst)}


def login(client, user):
    client.force_login(user)
    return client


def test_seeded_reconciliation_is_consistent(demo):
    inst, _ = demo
    recon = Reconciliation.objects.get(institution=inst)
    stmt = build_statement(recon)
    assert stmt.left.balance_check_difference == 0 and stmt.right.balance_check_difference == 0
    assert stmt.final_difference == 0
    assert recon.run_summary["auto_accepted"] == 4 and recon.run_summary["high_confidence"] == 1
    assert len(stmt.duplicates) == 1  # repeated SMS charge flagged


def test_all_pages_render_for_cfo(client, demo):
    inst, users = demo
    login(client, users["cfo"])
    recon = Reconciliation.objects.get(institution=inst)
    journal = Journal.objects.filter(institution=inst).first()
    batch = ImportBatch.objects.filter(institution=inst).first()
    item = SuspenseItem.objects.filter(institution=inst).first()
    account = Account.objects.get(institution=inst, code="1201")
    from apps.audit.models import Engagement, Finding

    urls = [
        reverse("reporting:dashboard"),
        reverse("accounting:accounts"),
        reverse("accounting:periods"),
        reverse("accounting:journals"),
        reverse("accounting:journal-detail", args=[journal.pk]),
        reverse("imports:batches"),
        reverse("imports:batch-detail", args=[batch.pk]),
        reverse("imports:profiles"),
        reverse("imports:transactions"),
        reverse("reconciliation:list"),
        reverse("reconciliation:rules"),
        reverse("settlements:dashboard"),
        reverse("suspense:list"),
        reverse("suspense:aging"),
        reverse("suspense:detail", args=[item.pk]),
        reverse("reporting:trial-balance"),
        reverse("reporting:general-ledger") + f"?account={account.pk}",
        reverse("reporting:income-statement"),
        reverse("reporting:financial-position"),
        reverse("reporting:recon-statement", args=[recon.pk]),
        reverse("reporting:suspense-aging"),
        reverse("audit:engagements"),
        reverse("audit:engagement", args=[Engagement.objects.get(institution=inst).pk]),
        reverse("audit:finding", args=[Finding.objects.get(institution=inst).pk]),
        reverse("organizations:branches"),
        reverse("notifications:list"),
        reverse("accounts:profile"),
    ] + [
        reverse("reconciliation:detail", args=[recon.pk]) + f"?tab={t}"
        for t in ("exceptions", "review", "matches", "statement", "workflow")
    ]
    for url in urls:
        resp = client.get(url)
        assert resp.status_code == 200, url


def test_admin_pages(client, demo):
    _, users = demo
    login(client, users["admin"])
    for name in (
        "organizations:settings",
        "accounts:users",
        "accounts:roles",
        "accounts:user-create",
        "accounts:role-create",
        "activity-log",
        "imports:profile-create",
        "reconciliation:rule-create",
    ):
        assert client.get(reverse(name)).status_code == 200, name


def test_role_based_navigation(client, demo):
    _, users = demo
    login(client, users["viewer"])
    html = client.get(reverse("reporting:dashboard")).content.decode()
    assert "Chart of accounts" not in html and "Trial balance" in html
    login(client, users["extaudit"])
    assert client.get(reverse("accounting:journal-create")).status_code == 403
    assert client.get(reverse("accounting:journals")).status_code == 200


def test_journal_workflow_through_ui(client, demo):
    inst, users = demo
    rent = Account.objects.get(institution=inst, code="5601")
    accrual = Account.objects.get(institution=inst, code="2301")
    from django.utils import timezone

    today = timezone.localdate().isoformat()
    login(client, users["accountant"])
    data = {
        "journal_type": "ACCRUAL",
        "transaction_date": today,
        "accounting_date": today,
        "description": "UI accrual",
        "currency": "NGN",
        "exchange_rate": "1",
        "form-TOTAL_FORMS": "4",
        "form-INITIAL_FORMS": "0",
        "form-MIN_NUM_FORMS": "0",
        "form-MAX_NUM_FORMS": "1000",
        "form-0-account": rent.pk,
        "form-0-debit": "1000.00",
        "form-1-account": accrual.pk,
        "form-1-credit": "1000.00",
    }
    resp = client.post(reverse("accounting:journal-create"), data)
    assert resp.status_code == 302
    journal = Journal.objects.get(institution=inst, description="UI accrual")
    client.post(reverse("accounting:journal-action", args=[journal.pk]), {"action": "submit"})
    login(client, users["finance"])
    client.post(reverse("accounting:journal-action", args=[journal.pk]), {"action": "approve"})
    client.post(reverse("accounting:journal-action", args=[journal.pk]), {"action": "post"})
    journal.refresh_from_db()
    assert journal.status == Journal.Status.POSTED


def test_unbalanced_journal_form_shows_error(client, demo):
    inst, users = demo
    rent = Account.objects.get(institution=inst, code="5601")
    accrual = Account.objects.get(institution=inst, code="2301")
    login(client, users["accountant"])
    data = {
        "journal_type": "MANUAL",
        "transaction_date": "2026-10-01",
        "accounting_date": "2026-10-01",
        "description": "Unbalanced",
        "currency": "NGN",
        "exchange_rate": "1",
        "form-TOTAL_FORMS": "2",
        "form-INITIAL_FORMS": "0",
        "form-MIN_NUM_FORMS": "0",
        "form-MAX_NUM_FORMS": "1000",
        "form-0-account": rent.pk,
        "form-0-debit": "1000.00",
        "form-1-account": accrual.pk,
        "form-1-credit": "900.00",
    }
    client.post(reverse("accounting:journal-create"), data)
    journal = Journal.objects.get(description="Unbalanced")
    resp = client.post(reverse("accounting:journal-action", args=[journal.pk]), {"action": "submit"}, follow=True)
    assert "not balanced" in resp.content.decode()
    journal.refresh_from_db()
    assert journal.status == Journal.Status.DRAFT


def test_import_upload_through_ui(client, demo):
    inst, users = demo
    profile = ImportProfile.objects.get(institution=inst, source_type="BANK_STATEMENT")
    login(client, users["recon"])
    content = b"Txn Date,Value Date,Reference,Description,Debit,Credit\n01/09/2026,01/09/2026,UI-1,Test,,100.00\n"
    upload = io.BytesIO(content)
    upload.name = "ui.csv"
    resp = client.post(reverse("imports:upload"), {"profile": profile.pk, "file": upload})
    assert resp.status_code == 302
    batch = ImportBatch.objects.filter(institution=inst, original_filename="ui.csv").get()
    assert "Preview" in client.get(reverse("imports:batch-detail", args=[batch.pk])).content.decode()
    client.post(reverse("imports:batch-action", args=[batch.pk]), {"action": "validate"})
    client.post(reverse("imports:batch-action", args=[batch.pk]), {"action": "commit"})
    batch.refresh_from_db()
    assert batch.status == ImportBatch.Status.COMMITTED and batch.committed_rows == 1


def test_reconciliation_review_and_approval_through_ui(client, demo):
    inst, users = demo
    recon = Reconciliation.objects.get(institution=inst)
    url = reverse("reconciliation:action", args=[recon.pk])
    login(client, users["recon"])
    group = recon.match_groups.get(status=MatchGroup.Status.PROPOSED)
    client.post(url, {"action": "accept", "group": group.pk, "note": "Confirmed with customer"})
    client.post(url, {"action": "submit", "note": "Ready"})
    recon.refresh_from_db()
    assert recon.status == Reconciliation.Status.SUBMITTED
    resp = client.post(url, {"action": "approve"}, follow=True)
    assert resp.status_code == 403 or "permission" in resp.content.decode().lower()
    login(client, users["reviewer"])
    client.post(url, {"action": "approve", "note": "Reviewed"})
    recon.refresh_from_db()
    assert recon.status == Reconciliation.Status.APPROVED


def test_exports(client, demo):
    inst, users = demo
    login(client, users["cfo"])
    for name in ("reporting:trial-balance", "reporting:income-statement", "reporting:financial-position"):
        for fmt, ctype in (("xlsx", "spreadsheetml"), ("csv", "text/csv"), ("pdf", "application/pdf")):
            resp = client.get(reverse(name) + f"?export={fmt}")
            assert resp.status_code == 200 and ctype in resp["Content-Type"], (name, fmt)
    resp = client.get(reverse("reporting:trial-balance") + "?export=xlsx")
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws["A2"].value == "Trial Balance"
    values = [row for row in ws.iter_rows(values_only=True) if row and row[1] == "Totals"]
    assert values and Decimal(str(values[0][5])) == Decimal(str(values[0][6]))
    recon = Reconciliation.objects.get(institution=inst)
    resp = client.get(reverse("reporting:recon-statement", args=[recon.pk]) + "?export=pdf")
    assert resp.content[:4] == b"%PDF"
    assert client.get(reverse("reporting:trial-balance") + "?export=exe").status_code == 403


def test_formula_injection_is_neutralised():
    assert safe_text('=HYPERLINK("http://evil")') == '\'=HYPERLINK("http://evil")'
    assert safe_text("@SUM(A1)") == "'@SUM(A1)"
    assert safe_text("-1,250.00") == "-1,250.00"
    assert safe_text("Normal text") == "Normal text"


def test_evidence_upload_and_download(client, demo):
    inst, users = demo
    journal = Journal.objects.filter(institution=inst).first()
    login(client, users["accountant"])
    f = io.BytesIO(b"%PDF-1.4 evidence")
    f.name = "invoice.pdf"
    resp = client.post(
        reverse("evidence-upload", args=["journal", journal.pk]), {"file": f, "description": "Invoice", "next": "/"}
    )
    assert resp.status_code == 302
    from apps.core.models import Attachment

    att = Attachment.objects.get(institution=inst, object_id=str(journal.pk))
    assert len(att.sha256) == 64
    assert client.get(reverse("evidence-download", args=[att.pk])).status_code == 200
    bad = io.BytesIO(b"MZ")
    bad.name = "virus.exe"
    client.post(reverse("evidence-upload", args=["journal", journal.pk]), {"file": bad, "next": "/"})
    assert Attachment.objects.filter(institution=inst).count() == 1
